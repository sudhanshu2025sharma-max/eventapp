import io, json
from zoneinfo import ZoneInfo
from django.shortcuts import render, redirect
from django.http import JsonResponse
from django.contrib.auth.decorators import login_required
from django.views.decorators.csrf import csrf_exempt
from django.db import models
from django.utils import timezone
from django.contrib import messages
from django.conf import settings
from django.core.mail import EmailMessage

from apps.accounts.models import User
from apps.accounts.permissions_helper import module_required
from apps.checkins.models import CheckIn, MealPass, MealWindow
from apps.checkins.meal_utils import sync_meal_window
from apps.checkins.push import push_to_checked_in
from apps.leaderboard.utils import award_points

IST = ZoneInfo("Asia/Kolkata")


def _today():
    return timezone.now().astimezone(IST).date()


def _meal_payload(window):
    if not window:
        return {"is_open": False, "meal_type": "Lunch", "start_time": None, "end_time": None, "scanned_count": 0}
    today = _today()
    scanned_count = MealPass.objects.filter(date=today, meal_type=window.meal_type, used=True).count()
    return {
        "is_open": bool(window.is_open),
        "meal_type": window.meal_type or "Lunch",
        "start_time": window.start_time.strftime("%I:%M %p") if window.start_time else None,
        "end_time": window.end_time.strftime("%I:%M %p") if window.end_time else None,
        "date": str(window.date),
        "scanned_count": scanned_count,
    }


@login_required
@module_required('checkin_scanner')
def scanner_view(request):
    try:
        sync_meal_window()
    except Exception:
        pass
    today = _today()
    window = MealWindow.objects.filter(date=today).order_by('-id').first()
    return render(request, 'panel/scanner.html', {
        "is_open": bool(window and window.is_open),
        "meal_window": window,
        "meal_type": window.meal_type if window else "Lunch",
        "today": today,
    })


@login_required
@module_required('checkin_scanner')
def checkin_list_view(request):
    today = _today()
    tab = request.GET.get('tab', 'checked_in')
    search = request.GET.get('search', '').strip()
    selected_meal = request.GET.get('meal_type', '').strip()

    tabs = [
        ('checked_in', 'Checked In Attendees'),
        ('not_checked_in', 'Not Checked In'),
        ('meal', 'Meal Pass History'),
    ]

    checked_in_qs = CheckIn.objects.select_related('user', 'scanned_by').filter(checkin_type='conference').order_by('-scanned_at')
    
    if search:
        checked_in_qs = checked_in_qs.filter(
            models.Q(user__first_name__icontains=search) |
            models.Q(user__last_name__icontains=search) |
            models.Q(user__email__icontains=search) |
            models.Q(user__registration_id__icontains=search)
        )

    checked_in_count = CheckIn.objects.filter(checkin_type='conference').count()
    total_users_count = User.objects.filter(
        models.Q(role='participant') | models.Q(role='attendee') | models.Q(role='speaker')
    ).count() or User.objects.count()
    not_checked_in_count = max(0, total_users_count - checked_in_count)

    meal_passes_qs = MealPass.objects.select_related('user', 'scanned_by').filter(date=today).order_by('-used_at', '-created_at')
    
    if selected_meal:
        meal_passes_qs = meal_passes_qs.filter(meal_type__iexact=selected_meal)
        
    if search:
        meal_passes_qs = meal_passes_qs.filter(
            models.Q(user__first_name__icontains=search) |
            models.Q(user__last_name__icontains=search) |
            models.Q(user__email__icontains=search) |
            models.Q(guest_name__icontains=search) |
            models.Q(guest_email__icontains=search)
        )

    checked_in_user_ids = list(CheckIn.objects.filter(checkin_type='conference').values_list('user_id', flat=True))
    not_checked_in_users = User.objects.filter(
        models.Q(role='participant') | models.Q(role='attendee') | models.Q(role='speaker'),
        is_active=True
    ).exclude(id__in=checked_in_user_ids).order_by('first_name')

    if search and tab == 'not_checked_in':
        not_checked_in_users = not_checked_in_users.filter(
            models.Q(first_name__icontains=search) |
            models.Q(last_name__icontains=search) |
            models.Q(email__icontains=search) |
            models.Q(registration_id__icontains=search)
        )

    # Unique meal categories available today for filter dropdown
    available_meals = list(MealPass.objects.filter(date=today).values_list('meal_type', flat=True).distinct())

    return render(request, 'panel/checkin_list.html', {
        "today": today,
        "tab": tab,
        "tabs": tabs,
        "search": search,
        "selected_meal": selected_meal,
        "available_meals": available_meals,
        "checkins": checked_in_qs[:200],
        "not_checked_in": not_checked_in_users[:200],
        "meal_passes": meal_passes_qs[:200],
        "checked_in_count": checked_in_count,
        "not_checked_in_count": not_checked_in_count,
        "meal_used_count": MealPass.objects.filter(date=today, used=True).count(),
    })


@csrf_exempt
@login_required
@module_required('checkin_scanner')
def panel_scan(request):
    if request.method != 'POST':
        return JsonResponse({"success": False, "message": "POST required"}, status=405)
    try:
        data = json.loads(request.body.decode('utf-8'))
    except Exception:
        data = request.POST
    code = str(data.get('registration_id') or data.get('qr_data') or data.get('qr_code') or '').strip()
    user_id = data.get('user_id')

    target_user = None
    if user_id:
        target_user = User.objects.filter(id=user_id).first()
    elif len(code) == 36 and '-' in code:
        target_user = User.objects.filter(id=code).first()

    if not target_user:
        users = list(User.objects.filter(models.Q(registration_id__iexact=code) | models.Q(email__iexact=code)))
        if len(users) > 1:
            return JsonResponse({
                "success": True,
                "multiple": True,
                "users": [{
                    "id": str(u.id),
                    "name": u.get_full_name() or u.email,
                    "email": u.email,
                    "registration_id": getattr(u, 'registration_id', '') or '',
                    "category": getattr(u, 'category', '') or u.get_role_display(),
                    "temp_password": getattr(u, 'temp_password', '') if getattr(u, 'must_change_password', False) else "Changed"
                } for u in users]
            })
        target_user = users[0] if users else None
    if not target_user:
        return JsonResponse({"success": False, "message": "Participant not found"}, status=404)
    checkin, created = CheckIn.objects.get_or_create(
        user=target_user, checkin_type='conference',
        defaults={'scanned_by': request.user},
    )
    return JsonResponse({
        "success": True,
        "already_checked_in": not created,
        "message": f"{target_user.get_full_name() or target_user.email} checked in successfully!",
        "user": {
            "name": target_user.get_full_name() or target_user.email,
            "email": target_user.email,
            "registration_id": getattr(target_user, 'registration_id', '') or '',
            "goodies_status": checkin.goodies_status,
        },
    })


@csrf_exempt
@login_required
@module_required('checkin_scanner')
def panel_goodies(request):
    if request.method != 'POST':
        return JsonResponse({"success": False, "message": "POST required"}, status=405)
    try:
        data = json.loads(request.body.decode('utf-8'))
    except Exception:
        data = request.POST
    code = str(data.get('registration_id') or data.get('qr_data') or data.get('user_id') or '').strip()
    action = str(data.get('action') or 'received').strip()
    note = str(data.get('note') or '').strip()

    target_user = User.objects.filter(
        models.Q(registration_id__iexact=code) |
        models.Q(email__iexact=code) |
        models.Q(id=code if len(code)==36 else None)
    ).first()

    if not target_user:
        return JsonResponse({"success": False, "message": "Participant not found"}, status=404)

    checkin, _ = CheckIn.objects.get_or_create(
        user=target_user, checkin_type='conference',
        defaults={'scanned_by': request.user},
    )
    checkin.goodies_status = action
    checkin.goodies_note = note
    checkin.goodies_confirmed_by = request.user
    checkin.goodies_confirmed_at = timezone.now()
    checkin.save()
    return JsonResponse({
        "success": True,
        "message": f"Kit status updated for {target_user.get_full_name() or target_user.email}!",
        "goodies_status": action,
        "goodies_note": note,
    })


@login_required
@module_required('meal_scanner')
def panel_meal_window_status(request):
    try:
        sync_meal_window()
    except Exception:
        pass
    window = MealWindow.objects.filter(date=_today()).order_by('-id').first()
    meal = _meal_payload(window)
    return JsonResponse({
        "success": True,
        "meal": meal,
        "is_open": meal["is_open"],
        "meal_open": meal["is_open"],
        "meal_type": meal["meal_type"],
        "meal_name": meal["meal_type"],
        "start_time": meal.get("start_time"),
        "end_time": meal.get("end_time"),
    })


@csrf_exempt
@login_required
@module_required('meal_scanner')
def panel_meal_window_toggle(request):
    action = None
    meal_type = 'Lunch'
    try:
        body = json.loads(request.body.decode('utf-8'))
        action = body.get('action')
        meal_type = (body.get('meal_type') or body.get('meal_name') or 'Lunch').strip()
    except Exception:
        action = request.POST.get('action') or request.GET.get('action')
        meal_type = (request.POST.get('meal_type') or request.GET.get('meal_type') or 'Lunch').strip()

    target_state = (str(action).lower() == 'open')
    today = _today()
    window = MealWindow.objects.filter(date=today).order_by('-id').first()
    if not window:
        window = MealWindow.objects.create(date=today, is_open=target_state, meal_type=meal_type)

    window.is_open = target_state
    if target_state:
        window.meal_type = meal_type
    window.opened_by = request.user
    if not target_state:
        window.closed_at = timezone.now()
    window.save()

    active_meal_name = window.meal_type or "Meal"

    if target_state:
        push_to_checked_in(
            title=f"🍽️ {active_meal_name} service is now open!",
            body=f"Dining service for {active_meal_name} is active.",
            data={"type": "meal_pass", "screen": "qr"},
        )
    else:
        push_to_checked_in(
            title=f"🛑 {active_meal_name} service is now closed",
            body=f"Dining service for {active_meal_name} has concluded.",
            data={"type": "meal_pass", "screen": "qr"},
        )
    return JsonResponse({
        "success": True,
        "meal": _meal_payload(window),
        "is_open": window.is_open,
        "message": f"{active_meal_name} window opened" if target_state else "Window closed",
    })


@csrf_exempt
@login_required
@module_required('meal_scanner')
def panel_meal_push(request):
    """Re-broadcast push notification for active meal service."""
    today = _today()
    window = MealWindow.objects.filter(date=today).order_by('-id').first()
    if not window or not window.is_open:
        return JsonResponse({"success": False, "message": "No active meal window is open."}, status=400)

    meal_name = window.meal_type or "Meal"
    push_to_checked_in(
        title=f"🍽️ Reminder: {meal_name} is open!",
        body=f"Dining service for {meal_name} is currently open.",
        data={"type": "meal_pass", "screen": "qr"}
    )
    return JsonResponse({"success": True, "message": f"Push notification sent for {meal_name}!"})


@csrf_exempt
@login_required
@module_required('meal_scanner')
def panel_meal_window_schedule(request):
    return JsonResponse({"success": True})


panel_meal_schedule = panel_meal_window_schedule


@csrf_exempt
@login_required
@module_required('meal_scanner')
def panel_meal_scan(request):
    if request.method != 'POST':
        return JsonResponse({"success": False, "message": "POST required"}, status=405)
    import json
    try:
        data = json.loads(request.body.decode('utf-8'))
    except Exception:
        data = request.POST
    
    raw_qr = str(data.get('qr_data') or data.get('qr_code') or data.get('registration_id') or '').strip()
    
    user_id = data.get("user_id")
    if user_id:
        target_user = User.objects.filter(id=user_id).first()
    else:
        users = list(User.objects.filter(models.Q(registration_id__iexact=raw_qr) | models.Q(email__iexact=raw_qr)))
        if not users:
            try:
                import uuid
                fallback_pass = MealPass.objects.filter(id=uuid.UUID(raw_qr)).first()
                if fallback_pass and fallback_pass.user: users = [fallback_pass.user]
            except Exception: pass
        if len(users) > 1:
            return JsonResponse({"success": True, "multiple": True, "users": [{"id": str(u.id), "name": u.get_full_name() or u.email, "email": u.email, "category": getattr(u, "category", "") or getattr(u, "role", ""), "temp_password": getattr(u, "temp_password", "") if getattr(u, "must_change_password", False) else "Changed"} for u in users]})
        target_user = users[0] if users else None
    
    if not target_user:
        try:
            import uuid
            pass_uuid = uuid.UUID(raw_qr)
            fallback_pass = MealPass.objects.filter(id=pass_uuid).first()
            if fallback_pass and fallback_pass.user:
                target_user = fallback_pass.user
        except Exception:
            pass

    if not target_user:
        return JsonResponse({"success": False, "message": "Participant not found. Invalid ID."}, status=404)

    is_checked_in = CheckIn.objects.filter(user=target_user, checkin_type='conference').exists()
    if not is_checked_in:
        return JsonResponse({"success": False, "message": "Not Checked-In! User must check into the conference first."}, status=400)

    today = timezone.now().astimezone(IST).date()
    window = MealWindow.objects.filter(date=today).order_by('-id').first()
    if not window or not window.is_open:
        return JsonResponse({"success": False, "message": "No active dining service is currently open."}, status=400)

    active_meal_name = window.meal_type or "Meal"

    existing_scan = MealPass.objects.filter(user=target_user, date=today, meal_type=active_meal_name, used=True).first()
    if existing_scan:
        time_str = existing_scan.used_at.astimezone(IST).strftime('%I:%M %p') if existing_scan.used_at else ""
        return JsonResponse({"success": False, "message": f"Already Redeemed at {time_str}"}, status=400)

    meal_pass = MealPass.objects.filter(user=target_user, date=today, meal_type=active_meal_name, used=False).first()
    if not meal_pass:
        meal_pass = MealPass(user=target_user, date=today, meal_type=active_meal_name, is_active=True)
    
    meal_pass.used = True
    meal_pass.used_at = timezone.now()
    meal_pass.scanned_by = request.user
    meal_pass.save()

    return JsonResponse({
        "success": True,
        "message": f"✅ {active_meal_name} approved for {target_user.get_full_name() or target_user.email}!"
    })



@login_required
@module_required('checkin_scanner')
def panel_stats(request):
    today = _today()
    total = User.objects.filter(
        models.Q(role='participant') | models.Q(role='attendee') | models.Q(role='speaker')
    ).count() or User.objects.count()
    checked_in = CheckIn.objects.filter(checkin_type='conference').count()
    remaining = max(0, total - checked_in)
    return JsonResponse({
        "success": True,
        "checked_in": checked_in,
        "remaining": remaining,
        "total": total,
        "meals_served": MealPass.objects.filter(date=today, used=True).count(),
    })


@csrf_exempt
@login_required
@module_required('meal_scanner')
def meal_pass_create_view(request):
    if request.method != 'POST':
        return render(request, 'panel/meal_pass_manage.html')
    guest_name = (request.POST.get('guest_name') or '').strip()
    guest_email = (request.POST.get('guest_email') or '').strip()
    guest_phone = (request.POST.get('guest_phone') or '').strip()
    guest_reg_no = (request.POST.get('guest_reg_no') or '').strip()
    meal_type = (request.POST.get('meal_type') or 'Lunch').strip()
    if not guest_name:
        return JsonResponse({"success": False, "message": "Guest name required"}, status=400)
    mp = MealPass.objects.create(
        guest_name=guest_name,
        guest_email=guest_email,
        guest_phone=guest_phone,
        guest_reg_no=guest_reg_no,
        meal_type=meal_type,
        date=_today(),
        is_active=True,
    )
    email_sent = False
    if guest_email:
        try:
            import qrcode
            qr = qrcode.QRCode(version=1, box_size=10, border=4)
            qr.add_data(str(mp.id))
            qr.make(fit=True)
            img = qr.make_image(fill_color="black", back_color="white")
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            email = EmailMessage(
                subject=f"Your ETD 2026 {meal_type} Pass",
                body=f"Hello {guest_name},\n\nAttached is your {meal_type} pass QR code for ETD 2026.\n\nETD 2026 Organising Team",
                from_email=getattr(settings, 'DEFAULT_FROM_EMAIL', None),
                to=[guest_email],
            )
            email.attach(f"meal_pass_{mp.id}.png", buf.getvalue(), "image/png")
            email.send(fail_silently=True)
            email_sent = True
        except Exception as e:
            print("meal pass email failed:", e)
    ct = request.content_type or ''
    if 'application/x-www-form-urlencoded' in ct or request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        return JsonResponse({
            "success": True,
            "email_sent": email_sent,
            "pass": {"id": str(mp.id), "qr_data": str(mp.id), "display_name": mp.display_name},
        })
    messages.success(request, f"{meal_type} pass created for {guest_name}!")
    return redirect('/panel/checkins/scanner/')

import csv
from django.http import HttpResponse

@login_required
@module_required('checkin_scanner')
def export_checkins_csv(request):
    import csv
    from django.http import HttpResponse

    tab = (request.GET.get('tab') or 'checked_in').strip()
    search = (request.GET.get('search') or '').strip()
    selected_meal = (request.GET.get('meal_type') or '').strip()
    date_str = (request.GET.get('date') or '').strip()
    today = _today()

    # Optional date filter for meal history (YYYY-MM-DD)
    meal_date = today
    if date_str:
        try:
            from datetime import datetime as _dt
            meal_date = _dt.strptime(date_str, '%Y-%m-%d').date()
        except Exception:
            meal_date = today

    if tab == 'not_checked_in':
        filename = f"not_checked_in_{today}.csv"
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        writer = csv.writer(response)
        writer.writerow(['Registration ID', 'First Name', 'Last Name', 'Email', 'Role', 'Affiliation', 'Phone', 'Designation'])

        checked_in_ids = list(CheckIn.objects.filter(checkin_type='conference').values_list('user_id', flat=True))
        qs = User.objects.filter(
            models.Q(role='participant') | models.Q(role='attendee') | models.Q(role='speaker'),
            is_active=True
        ).exclude(id__in=checked_in_ids).order_by('first_name')

        if search:
            qs = qs.filter(
                models.Q(first_name__icontains=search) |
                models.Q(last_name__icontains=search) |
                models.Q(email__icontains=search) |
                models.Q(registration_id__icontains=search)
            )

        for u in qs:
            writer.writerow([
                getattr(u, 'registration_id', '') or '',
                u.first_name or '',
                u.last_name or '',
                u.email or '',
                getattr(u, 'role', '') or '',
                getattr(u, 'affiliation', '') or '',
                getattr(u, 'phone', '') or '',
                getattr(u, 'designation', '') or '',
            ])
        return response

    if tab == 'meal':
        filename = f"meal_history_{meal_date}.csv"
        response = HttpResponse(content_type='text/csv')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        writer = csv.writer(response)
        writer.writerow([
            'Date', 'Meal Type', 'Registration ID', 'Name', 'Email',
            'Status', 'Redeemed At (IST)', 'Scanned By', 'Pass ID', 'Guest Phone'
        ])

        qs = MealPass.objects.select_related('user', 'scanned_by').filter(date=meal_date).order_by('meal_type', '-used_at', '-created_at')
        if selected_meal:
            qs = qs.filter(meal_type__iexact=selected_meal)
        if search:
            qs = qs.filter(
                models.Q(user__first_name__icontains=search) |
                models.Q(user__last_name__icontains=search) |
                models.Q(user__email__icontains=search) |
                models.Q(user__registration_id__icontains=search) |
                models.Q(guest_name__icontains=search) |
                models.Q(guest_email__icontains=search)
            )

        for p in qs:
            redeemed = ''
            if p.used_at:
                redeemed = p.used_at.astimezone(IST).strftime('%Y-%m-%d %H:%M')
            writer.writerow([
                str(p.date),
                p.meal_type or 'Meal',
                p.display_reg or '',
                p.display_name or '',
                p.display_email or '',
                'Redeemed' if p.used else 'Active/Unused',
                redeemed,
                p.scanned_by.get_full_name() if p.scanned_by else '',
                str(p.id),
                p.guest_phone or '',
            ])
        return response

    # Default: checked_in
    filename = f"checked_in_{today}.csv"
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    writer = csv.writer(response)
    writer.writerow([
        'Registration ID', 'First Name', 'Last Name', 'Email', 'Role',
        'Affiliation', 'Phone', 'Kit Status', 'Kit Note', 'Scanned By', 'Checked In At (IST)'
    ])

    qs = CheckIn.objects.select_related('user', 'scanned_by').filter(checkin_type='conference').order_by('-scanned_at')
    if search:
        qs = qs.filter(
            models.Q(user__first_name__icontains=search) |
            models.Q(user__last_name__icontains=search) |
            models.Q(user__email__icontains=search) |
            models.Q(user__registration_id__icontains=search)
        )

    for c in qs:
        u = c.user
        scanned_str = c.scanned_at.astimezone(IST).strftime('%Y-%m-%d %H:%M') if c.scanned_at else ''
        writer.writerow([
            getattr(u, 'registration_id', '') or '',
            u.first_name or '',
            u.last_name or '',
            u.email or '',
            getattr(u, 'role', '') or '',
            getattr(u, 'affiliation', '') or '',
            getattr(u, 'phone', '') or '',
            c.goodies_status or 'pending',
            c.goodies_note or '',
            c.scanned_by.get_full_name() if c.scanned_by else '',
            scanned_str,
        ])
    return response


@login_required
@module_required('checkin_scanner')
def export_checkins_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from django.http import HttpResponse

    wb = openpyxl.Workbook()
    
    # ── SHEET 1: Checked In Attendees ──
    ws1 = wb.active
    ws1.title = "Checked In Attendees"
    ws1.append(['Registration ID', 'First Name', 'Last Name', 'Email', 'Role', 'Affiliation', 'Phone', 'Kit Status', 'Kit Note', 'Scanned By', 'Checked In At (IST)'])
    
    checkins = CheckIn.objects.select_related('user', 'scanned_by').filter(checkin_type='conference').order_by('-scanned_at')
    for c in checkins:
        u = c.user
        scanned_str = c.scanned_at.astimezone(IST).strftime('%Y-%m-%d %H:%M') if c.scanned_at else ''
        ws1.append([
            getattr(u, 'registration_id', '') or '',
            u.first_name or '',
            u.last_name or '',
            u.email or '',
            getattr(u, 'role', '') or '',
            getattr(u, 'affiliation', '') or '',
            getattr(u, 'phone', '') or '',
            c.goodies_status or 'pending',
            c.goodies_note or '',
            c.scanned_by.get_full_name() if c.scanned_by else '',
            scanned_str
        ])

    # ── SHEET 2: Not Checked In ──
    ws2 = wb.create_sheet(title="Not Checked In")
    ws2.append(['Registration ID', 'First Name', 'Last Name', 'Email', 'Role', 'Affiliation', 'Phone', 'Designation'])
    
    checked_in_ids = list(CheckIn.objects.filter(checkin_type='conference').values_list('user_id', flat=True))
    not_checked_in = User.objects.filter(
        models.Q(role='participant') | models.Q(role='attendee') | models.Q(role='speaker'),
        is_active=True
    ).exclude(id__in=checked_in_ids).order_by('first_name')

    for u in not_checked_in:
        ws2.append([
            getattr(u, 'registration_id', '') or '',
            u.first_name or '',
            u.last_name or '',
            u.email or '',
            getattr(u, 'role', '') or '',
            getattr(u, 'affiliation', '') or '',
            getattr(u, 'phone', '') or '',
            getattr(u, 'designation', '') or ''
        ])

    # ── SHEET 3: Meal Pass History ──
    ws3 = wb.create_sheet(title="Meal Pass History")
    ws3.append(['Date', 'Meal Type', 'Registration ID', 'Name', 'Email', 'Status', 'Redeemed At (IST)', 'Scanned By', 'Pass ID', 'Guest Phone'])
    
    meal_passes = MealPass.objects.select_related('user', 'scanned_by').order_by('-date', 'meal_type', '-used_at')
    for p in meal_passes:
        redeemed = p.used_at.astimezone(IST).strftime('%Y-%m-%d %H:%M') if p.used_at else ''
        ws3.append([
            str(p.date),
            p.meal_type or 'Meal',
            p.display_reg or '',
            p.display_name or '',
            p.display_email or '',
            'Redeemed' if p.used else 'Active/Unused',
            redeemed,
            p.scanned_by.get_full_name() if p.scanned_by else '',
            str(p.id),
            p.guest_phone or ''
        ])

    # Styling Headers & Column Widths
    header_fill = PatternFill(start_color="0333B6", end_color="0333B6", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="ETD2026_CheckIn_And_Meal_Report_{_today()}.xlsx"'
    wb.save(response)
    return response
